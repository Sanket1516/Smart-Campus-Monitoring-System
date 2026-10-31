import json, sys
from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_ORIENT
d = json.load(open(sys.argv[1]))
doc = Document()
s = doc.sections[0]; s.orientation = WD_ORIENT.LANDSCAPE
s.page_width, s.page_height = s.page_height, s.page_width
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10)
doc.add_heading('Smart Campus Monitor - Demo Credentials', 0)
doc.add_paragraph('Login URL: http://localhost:5173    |    API: http://localhost:5000')
def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Light Grid Accent 1'
    for c, h in zip(t.rows[0].cells, headers):
        c.text = h
        for r in c.paragraphs[0].runs: r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row): c.text = str(v)
doc.add_heading('1. Staff accounts (Admin, Warden, Security)', 1)
table(['Role','Username','Password','Name','Hostel','Email (receives alerts)','Phone'],
  [[a['role'].title(),a['username'],a['password'],a['name'],a['hostel'] or '-',a['email'],a['phone']] for a in d['accounts']])
doc.add_heading('2. Students and parent emails', 1)
doc.add_paragraph('Parent emails are routed to the four demo inboxes: ' + ', '.join(d['inboxes']) + '.')
table(['SAP ID','Name','Type','Course / Dept','Yr','Hostel / Room','Phone','Student email','Parent email','Parent phone'],
  [[x['sapId'],x['name'],x['category'],f"{x['course']} / {x['department']}",x['year'],f"{x['hostel']} / {x['room']}",x['phone'],x['email'],x['parentEmail'],x['parentPhone']] for x in d['students']])
doc.add_heading('3. Student exit-request portal (hostellers only)', 1)
doc.add_paragraph('URL: /student/exit-request. Username = student full name; Password = student phone number (see table above). Day scholars cannot log in.')
doc.add_heading('Notes', 1)
doc.add_paragraph('Demo data only. Change all passwords before any real deployment.')
doc.save(sys.argv[2])
